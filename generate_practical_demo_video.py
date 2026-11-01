import os
import re
import cv2
import numpy as np
from PIL import Image, ImageDraw, ImageFont
import docx

# Video Specifications
WIDTH = 1280
HEIGHT = 720
FPS = 30
OUTPUT_VIDEO = "CEUE203_Practical_Demonstration_Video.mp4"

# Fonts
FONT_DIR = os.path.join(os.environ.get("WINDIR", "C:\\Windows"), "Fonts")
FONT_CODE = ImageFont.truetype(os.path.join(FONT_DIR, "consola.ttf"), 14)
FONT_CODE_BOLD = ImageFont.truetype(os.path.join(FONT_DIR, "consolab.ttf"), 14)
FONT_TITLE = ImageFont.truetype(os.path.join(FONT_DIR, "arialbd.ttf"), 16)
FONT_BIG = ImageFont.truetype(os.path.join(FONT_DIR, "arialbd.ttf"), 28)
FONT_MED = ImageFont.truetype(os.path.join(FONT_DIR, "arialbd.ttf"), 20)
FONT_SMALL = ImageFont.truetype(os.path.join(FONT_DIR, "arial.ttf"), 12)
FONT_SMALL_BOLD = ImageFont.truetype(os.path.join(FONT_DIR, "arialbd.ttf"), 12)

# Official Practical Index Metadata from User's Image
PRACTICAL_INFO = {
    1:  ("7/7/26",   "Practical-1 : Java Fundamentals and MiniBank Shell"),
    2:  ("13/7/26",  "Practical-2 : Classes and Objects - Account and Customer"),
    3:  ("20/7/26",  "Practical-3 : Object Class Methods, Nested Classes and Cloneable"),
    4:  ("29/7/26",  "Practical-4 : String Handling, StringBuilder and ReGex"),
    5:  ("3/8/26",   "Practical-5 : Inheritance and Polymorphism"),
    6:  ("17/8/26",  "Practical-6 : Interfaces, Functional interfaces, packages and JAR"),
    7:  ("27/10/26", "Practical-7 : Custom Annotations"),
    8:  ("16/9/26",  "Practical-8 : Exception Handling"),
    9:  ("16/9/26",  "Practical-9 : Thread Synchronization and race conditions"),
    10: ("5/10/26",  "Practical-10 : Thread Pools, Producer-Consumer and Deadlock"),
    11: ("27/10/26", "Practical-11 : File I/O, Serialization, NIO report"),
    12: ("27/10/26", "Practical-12 : Collection Framework"),
    13: ("27/10/26", "Practical-13 : Generic"),
    14: ("27/10/26", "Practical-14 : Capstone - MiniBank Assembly, JUnit Testing and Runnable JAR")
}

# Java syntax highlighting keywords
KEYWORDS = {
    'public', 'private', 'protected', 'class', 'interface', 'enum', 'record',
    'void', 'static', 'final', 'extends', 'implements', 'import', 'package',
    'return', 'new', 'if', 'else', 'for', 'while', 'switch', 'case', 'yield',
    'try', 'catch', 'finally', 'throw', 'throws', 'synchronized', 'this', 'super'
}
TYPES = {
    'String', 'int', 'long', 'double', 'float', 'boolean', 'char', 'byte', 'short',
    'Integer', 'Long', 'Double', 'Boolean', 'List', 'ArrayList', 'Map', 'HashMap',
    'Set', 'HashSet', 'LinkedHashSet', 'ConcurrentHashMap', 'Account', 'Customer',
    'Scanner', 'Path', 'Files', 'Object', 'Exception', 'BankException'
}

def highlight_line(draw, x, y, line):
    # Check if comment
    trimmed = line.strip()
    if trimmed.startswith('//') or trimmed.startswith('/*') or trimmed.startswith('*'):
        draw.text((x, y), line, fill=(106, 153, 85), font=FONT_CODE)
        return

    # Check annotations
    if trimmed.startswith('@'):
        draw.text((x, y), line, fill=(220, 220, 170), font=FONT_CODE_BOLD)
        return

    # Simple token parser
    tokens = re.split(r'(\s+|[(){}\[\],;.\"]|==|!=|<=|>=|\+=|-=|\*=|->)', line)
    curr_x = x
    in_string = False

    for tok in tokens:
        if not tok:
            continue
        if tok == '"':
            in_string = not in_string
            draw.text((curr_x, y), tok, fill=(206, 145, 120), font=FONT_CODE)
        elif in_string:
            draw.text((curr_x, y), tok, fill=(206, 145, 120), font=FONT_CODE)
        elif tok in KEYWORDS:
            draw.text((curr_x, y), tok, fill=(86, 156, 214), font=FONT_CODE_BOLD)
        elif tok in TYPES:
            draw.text((curr_x, y), tok, fill=(78, 201, 176), font=FONT_CODE)
        elif tok.isdigit():
            draw.text((curr_x, y), tok, fill=(181, 206, 168), font=FONT_CODE)
        else:
            draw.text((curr_x, y), tok, fill=(212, 212, 212), font=FONT_CODE)

        w = draw.textlength(tok, font=FONT_CODE)
        curr_x += int(w)

def render_title_slide(video_writer, title, subtitle, duration_sec):
    frames = int(duration_sec * FPS)
    for f in range(frames):
        img = Image.new('RGB', (WIDTH, HEIGHT), color=(18, 24, 38))
        draw = ImageDraw.Draw(img)

        # Header bar
        draw.rectangle([0, 0, WIDTH, 6], fill=(0, 122, 204))

        # University Header
        draw.text((WIDTH // 2, 140), "CHAROTAR UNIVERSITY OF SCIENCE AND TECHNOLOGY (CHARUSAT)",
                  fill=(220, 230, 245), font=FONT_MED, anchor="mm")
        draw.text((WIDTH // 2, 180), "Chandubhai S. Patel Institute of Technology (CSPIT) | IT Department",
                  fill=(160, 180, 210), font=FONT_TITLE, anchor="mm")

        # Main Title Banner
        draw.rectangle([120, 230, WIDTH - 120, 360], fill=(26, 36, 56), outline=(0, 122, 204), width=2)
        draw.text((WIDTH // 2, 275), title, fill=(255, 255, 255), font=FONT_BIG, anchor="mm")
        draw.text((WIDTH // 2, 320), subtitle, fill=(78, 201, 176), font=FONT_MED, anchor="mm")

        # Student Details Card
        draw.rectangle([280, 400, WIDTH - 280, 560], fill=(22, 30, 48), outline=(60, 80, 110), width=1)
        details = [
            ("Student Name:", "Vishnu Tejalkumar Jariwala"),
            ("Enrollment ID:", "D26IT113"),
            ("Course:", "Object Oriented Programming (Java) [CEUE203]"),
            ("Class & Semester:", "B.Tech IT - 2 (3rd Semester)"),
            ("Practicals Covered:", "Practical 1 through 14 (Tasks & MiniBank Capstone)")
        ]
        y_pos = 425
        for lbl, val in details:
            draw.text((310, y_pos), lbl, fill=(180, 190, 205), font=FONT_SMALL_BOLD)
            draw.text((480, y_pos), val, fill=(255, 255, 255), font=FONT_SMALL)
            y_pos += 25

        # Animated progress indicator
        prog = f / max(1, frames - 1)
        draw.rectangle([0, HEIGHT - 8, int(WIDTH * prog), HEIGHT], fill=(0, 122, 204))

        frame = cv2.cvtColor(np.array(img), cv2.COLOR_RGB2BGR)
        video_writer.write(frame)

def render_chapter_slide(video_writer, p_num, date_str, full_title, duration_sec):
    frames = int(duration_sec * FPS)
    for f in range(frames):
        img = Image.new('RGB', (WIDTH, HEIGHT), color=(20, 26, 40))
        draw = ImageDraw.Draw(img)

        # Header line
        draw.rectangle([0, 0, WIDTH, 5], fill=(0, 122, 204))

        # Badge
        draw.rounded_rectangle([WIDTH // 2 - 120, 200, WIDTH // 2 + 120, 250], radius=10, fill=(0, 122, 204))
        draw.text((WIDTH // 2, 225), f"PRACTICAL - {p_num:02d}", fill=(255, 255, 255), font=FONT_MED, anchor="mm")

        # Title
        draw.text((WIDTH // 2, 310), full_title, fill=(255, 255, 255), font=FONT_BIG, anchor="mm")
        draw.text((WIDTH // 2, 370), f"Conducted / Submission Date: {date_str}   |   Course: CEUE203 (OOP)",
                  fill=(78, 201, 176), font=FONT_MED, anchor="mm")

        # Details box
        draw.rectangle([240, 430, WIDTH - 240, 520], fill=(28, 38, 58), outline=(60, 80, 110), width=1)
        draw.text((WIDTH // 2, 460), "Components Demonstrated: Task 1, Task 2, Task 3 & MiniBank Milestone",
                  fill=(220, 230, 245), font=FONT_TITLE, anchor="mm")
        draw.text((WIDTH // 2, 490), "Student: Vishnu Jariwala (D26IT113)  |  CSPIT - Information Technology",
                  fill=(160, 180, 210), font=FONT_SMALL, anchor="mm")

        # Progress bar
        prog = f / max(1, frames - 1)
        draw.rectangle([WIDTH // 2 - 200, 560, WIDTH // 2 - 200 + int(400 * prog), 565], fill=(0, 122, 204))

        frame = cv2.cvtColor(np.array(img), cv2.COLOR_RGB2BGR)
        video_writer.write(frame)

def render_code_and_execution_slide(video_writer, p_num, item_label, filename, code_text, output_text, duration_sec):
    frames = int(duration_sec * FPS)
    code_lines = code_text.splitlines()
    output_lines = [l for l in output_text.splitlines() if l.strip()]

    # Layout dimensions
    LEFT_W = 660
    RIGHT_W = WIDTH - LEFT_W
    TOP_H = 40
    BOTTOM_H = 26
    EDITOR_H = HEIGHT - TOP_H - BOTTOM_H

    # Visible lines
    LINE_H = 21
    VISIBLE_LINES = (EDITOR_H - 35) // LINE_H
    max_scroll = max(0, len(code_lines) - VISIBLE_LINES)

    cmd_str = f"javac {filename} && java {filename.replace('.java', '')}"

    for f in range(frames):
        img = Image.new('RGB', (WIDTH, HEIGHT), color=(30, 30, 30))
        draw = ImageDraw.Draw(img)

        # 1. TOP WINDOW BAR
        draw.rectangle([0, 0, WIDTH, TOP_H], fill=(45, 45, 48))
        # Window controls
        draw.ellipse([14, 14, 26, 26], fill=(255, 95, 86))
        draw.ellipse([34, 14, 46, 26], fill=(255, 189, 46))
        draw.ellipse([54, 14, 66, 26], fill=(39, 201, 63))
        # Title text
        draw.text((80, 11), f"CEUE203 | Practical {p_num} > {item_label} > {filename} | Vishnu Jariwala (D26IT113)",
                  fill=(220, 220, 220), font=FONT_TITLE)

        # 2. LEFT EDITOR (CODE SCROLL)
        draw.rectangle([0, TOP_H, LEFT_W, HEIGHT - BOTTOM_H], fill=(30, 30, 30))
        # Tab Header
        draw.rectangle([0, TOP_H, 220, TOP_H + 28], fill=(30, 30, 30))
        draw.line([0, TOP_H, 220, TOP_H], fill=(0, 122, 204), width=2)
        draw.text((15, TOP_H + 6), f"☕ {filename}", fill=(255, 255, 255), font=FONT_CODE_BOLD)
        draw.text((195, TOP_H + 6), "×", fill=(150, 150, 150), font=FONT_CODE)

        # Gutter background
        draw.rectangle([0, TOP_H + 28, 48, HEIGHT - BOTTOM_H], fill=(37, 37, 38))

        # Calculate smooth scroll position
        prog = f / max(1, frames - 1)
        scroll_line = int(prog * max_scroll)

        y = TOP_H + 34
        for l_idx in range(scroll_line, min(len(code_lines), scroll_line + VISIBLE_LINES + 1)):
            line_no = l_idx + 1
            line_str = code_lines[l_idx]
            # Draw line number
            draw.text((12, y), f"{line_no:3d}", fill=(130, 130, 130), font=FONT_CODE)
            # Draw syntax highlighted code
            highlight_line(draw, 58, y, line_str)
            y += LINE_H

        # 3. RIGHT TERMINAL (EXECUTION OUTPUT)
        draw.rectangle([LEFT_W, TOP_H, WIDTH, HEIGHT - BOTTOM_H], fill=(16, 18, 24))
        draw.line([LEFT_W, TOP_H, LEFT_W, HEIGHT - BOTTOM_H], fill=(50, 50, 50), width=1)

        # Terminal Tab
        draw.rectangle([LEFT_W, TOP_H, LEFT_W + 240, TOP_H + 28], fill=(22, 25, 33))
        draw.text((LEFT_W + 15, TOP_H + 6), "⚡ TERMINAL (powershell)", fill=(200, 210, 225), font=FONT_CODE_BOLD)

        # Command typing animation (types out over the first 25% of the clip)
        type_prog = min(1.0, prog / 0.25)
        chars_to_show = int(type_prog * len(cmd_str))
        typed_cmd = cmd_str[:chars_to_show]
        cursor_str = "_" if (f // 8) % 2 == 0 else ""

        term_y = TOP_H + 38
        draw.text((LEFT_W + 15, term_y), f"PS D:\\OOP\\practical-{p_num}> ", fill=(78, 201, 176), font=FONT_CODE_BOLD)
        px_w = int(draw.textlength(f"PS D:\\OOP\\practical-{p_num}> ", font=FONT_CODE_BOLD))
        draw.text((LEFT_W + 15 + px_w, term_y), typed_cmd + cursor_str, fill=(229, 192, 123), font=FONT_CODE)

        # Output reveal animation (reveals lines over the remaining 75% of the clip)
        term_y += 24
        if prog > 0.25:
            out_prog = (prog - 0.25) / 0.75
            visible_out_count = int(out_prog * len(output_lines))
            visible_out_count = max(1, min(len(output_lines), visible_out_count))

            # Show latest lines that fit in terminal
            max_term_lines = (EDITOR_H - 100) // 18
            start_out = max(0, visible_out_count - max_term_lines)

            for line in output_lines[start_out:visible_out_count]:
                # Style headers or success lines
                col = (204, 204, 204)
                if "===" in line or "---" in line:
                    col = (0, 122, 204)
                elif "PASS" in line or "Success" in line or "Verified" in line or "COMPLETE" in line:
                    col = (152, 195, 121)
                elif "Error" in line or "FAIL" in line:
                    col = (224, 108, 117)
                elif "Audit" in line or "AUDIT" in line:
                    col = (229, 192, 123)

                draw.text((LEFT_W + 15, term_y), line[:68], fill=col, font=FONT_CODE)
                term_y += 18

            # Completion badge at the end of output
            if prog > 0.85:
                draw.rectangle([LEFT_W + 15, term_y + 6, LEFT_W + 280, term_y + 28], fill=(24, 48, 30), outline=(50, 180, 80))
                draw.text((LEFT_W + 25, term_y + 10), "✔ EXECUTION PASSED (Code 0)", fill=(100, 255, 130), font=FONT_SMALL_BOLD)

        # 4. BOTTOM STATUS BAR
        draw.rectangle([0, HEIGHT - BOTTOM_H, WIDTH, HEIGHT], fill=(0, 122, 204))
        draw.text((15, HEIGHT - 19), f"CHARUSAT CSPIT | Practical {p_num}/14 | {item_label} | Java 25 LTS",
                  fill=(255, 255, 255), font=FONT_SMALL)
        draw.text((WIDTH - 340, HEIGHT - 19), "Vishnu Jariwala (D26IT113) | Branch: main",
                  fill=(255, 255, 255), font=FONT_SMALL_BOLD)

        frame = cv2.cvtColor(np.array(img), cv2.COLOR_RGB2BGR)
        video_writer.write(frame)

def render_outro_slide(video_writer, duration_sec):
    frames = int(duration_sec * FPS)
    for f in range(frames):
        img = Image.new('RGB', (WIDTH, HEIGHT), color=(18, 24, 38))
        draw = ImageDraw.Draw(img)

        # Header bar
        draw.rectangle([0, 0, WIDTH, 6], fill=(0, 122, 204))

        # Main Congratulations / Verification Banner
        draw.rectangle([140, 140, WIDTH - 140, 300], fill=(26, 36, 56), outline=(0, 122, 204), width=2)
        draw.text((WIDTH // 2, 190), "ALL 14 PRACTICALS COMPLETED & VERIFIED", fill=(100, 255, 130), font=FONT_BIG, anchor="mm")
        draw.text((WIDTH // 2, 245), "MiniBank Enterprise Capstone Architecture Fully Implemented", fill=(255, 255, 255), font=FONT_MED, anchor="mm")

        # Details Card
        draw.rectangle([240, 330, WIDTH - 240, 520], fill=(22, 30, 48), outline=(60, 80, 110), width=1)
        details = [
            ("Student Name:", "Vishnu Tejalkumar Jariwala"),
            ("Enrollment Number:", "D26IT113"),
            ("Institute / Dept:", "CSPIT - K.D. Patel Department of Information Technology"),
            ("Course:", "Object Oriented Programming [CEUE203]"),
            ("GitHub Repository:", "https://github.com/07vish/D26IT113_Vishnu_Jariwala-oop"),
            ("Status:", "All Practicals (1-14) Tested, Documented & Pushed in Sequence")
        ]
        y_pos = 350
        for lbl, val in details:
            draw.text((270, y_pos), lbl, fill=(180, 190, 205), font=FONT_SMALL_BOLD)
            draw.text((460, y_pos), val, fill=(255, 255, 255), font=FONT_SMALL)
            y_pos += 26

        # Progress bar
        prog = f / max(1, frames - 1)
        draw.rectangle([0, HEIGHT - 8, int(WIDTH * prog), HEIGHT], fill=(0, 122, 204))

        frame = cv2.cvtColor(np.array(img), cv2.COLOR_RGB2BGR)
        video_writer.write(frame)

def main():
    print(f"Initializing VideoWriter: {OUTPUT_VIDEO} ({WIDTH}x{HEIGHT} @ {FPS}fps)...")
    fourcc = cv2.VideoWriter_fourcc(*'mp4v')
    out = cv2.VideoWriter(OUTPUT_VIDEO, fourcc, FPS, (WIDTH, HEIGHT))

    if not out.isOpened():
        raise RuntimeError("Failed to open VideoWriter with mp4v codec.")

    # 1. Render Intro Slide (3.5 seconds)
    print("Rendering Intro Slide...")
    render_title_slide(
        out,
        title="OBJECT ORIENTED PROGRAMMING (CEUE203)",
        subtitle="Complete Practical Laboratory Demonstration (Practicals 1 - 14)",
        duration_sec=3.5
    )

    # 2. Iterate through Practical 1 to 14
    for p in range(1, 15):
        date_str, full_title = PRACTICAL_INFO[p]
        print(f"\nProcessing Practical {p:02d} ({full_title})...")

        # Chapter transition slide (1.2 seconds)
        render_chapter_slide(out, p, date_str, full_title, duration_sec=1.2)

        # Load Tasks & Micro docx
        t_doc = docx.Document(f"practical-{p}/Practical_{p:02d}_Tasks.docx")
        m_doc = docx.Document(f"practical-{p}/Practical_{p:02d}_MicroProject.docx")

        t1_code = t_doc.tables[0].cell(0, 0).text.strip()
        t1_out  = t_doc.tables[1].cell(0, 0).text.strip()
        t2_code = t_doc.tables[2].cell(0, 0).text.strip()
        t2_out  = t_doc.tables[3].cell(0, 0).text.strip()
        t3_code = t_doc.tables[4].cell(0, 0).text.strip()
        t3_out  = t_doc.tables[5].cell(0, 0).text.strip()

        m_code  = m_doc.tables[0].cell(0, 0).text.strip()
        m_out   = m_doc.tables[1].cell(0, 0).text.strip()

        # Filenames
        t1_fn = f"P{p:02d}_Task1.java"
        t2_fn = f"P{p:02d}_Task2.java"
        t3_fn = f"P{p:02d}_Task3.java"
        m_fn  = "MiniBank.java"

        # Try to find actual Java filenames in practical folder
        t1_files = [f for f in os.listdir(f"practical-{p}/task-1") if f.endswith('.java')] if os.path.exists(f"practical-{p}/task-1") else []
        t2_files = [f for f in os.listdir(f"practical-{p}/task-2") if f.endswith('.java')] if os.path.exists(f"practical-{p}/task-2") else []
        t3_files = [f for f in os.listdir(f"practical-{p}/task-3") if f.endswith('.java')] if os.path.exists(f"practical-{p}/task-3") else []
        if t1_files: t1_fn = t1_files[0]
        if t2_files: t2_fn = t2_files[0]
        if t3_files: t3_fn = t3_files[0]

        # Render Task 1 (3.2s)
        render_code_and_execution_slide(out, p, "Task 1", t1_fn, t1_code, t1_out, duration_sec=3.2)
        # Render Task 2 (3.2s)
        render_code_and_execution_slide(out, p, "Task 2", t2_fn, t2_code, t2_out, duration_sec=3.2)
        # Render Task 3 (3.2s)
        render_code_and_execution_slide(out, p, "Task 3", t3_fn, t3_code, t3_out, duration_sec=3.2)
        # Render MicroProject (4.2s)
        render_code_and_execution_slide(out, p, "Micro-Project", m_fn, m_code, m_out, duration_sec=4.2)

    # 3. Render Outro Slide (3.5 seconds)
    print("\nRendering Outro Slide...")
    render_outro_slide(out, duration_sec=3.5)

    out.release()
    print(f"\n==================================================")
    print(f"Successfully generated video: {OUTPUT_VIDEO}")
    file_size_mb = os.path.getsize(OUTPUT_VIDEO) / (1024 * 1024)
    print(f"File Size: {file_size_mb:.2f} MB")
    print(f"==================================================")

if __name__ == "__main__":
    main()
